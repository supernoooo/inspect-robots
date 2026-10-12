"""Upsert saved Colosseum scores into a local workbook without altering old sheets."""

from __future__ import annotations

import json
import math
import os
import shutil
import tempfile
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any

SHEET = "Colosseum Evals"
HEADERS = (
    "Router",
    "Evaluation ID",
    "Run ID",
    "Track",
    "Side",
    "Instruction",
    "Scene",
    "Evaluator",
    "Institution",
    "Model",
    "Model URL",
    "Revision",
    "Native control",
    "Success",
    "Score (0-1)",
    "Score (%)",
    "Preference",
    "Difficulty",
    "Feedback",
    "Test",
    "Result submitted",
    "Dataset uploaded",
    "Steps",
    "Duration (s)",
    "Recorded at",
    "Inspect log",
    "Dataset folder",
    "Head video",
    "Left video",
    "Right video",
)
_FILE_COLUMNS = ("Inspect log", "Dataset folder", "Head video", "Left video", "Right video")


class WorkbookBusy(RuntimeError):
    """The workbook is open in a spreadsheet editor or another exporter."""


def validate_target(value: str | Path) -> Path:
    """Validate format and optional dependencies before a robot evaluation starts."""
    path = Path(value).expanduser().resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Excel output must be an .xlsx file")
    try:
        import openpyxl

        del openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Install Excel support: uv pip install --python .venv/bin/python 'openpyxl>=3.1,<4'"
        ) from exc
    return path


def _read_optional(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected a JSON object in {path.name}")
    return value


def evaluation_rows(manifest: Path) -> list[dict[str, Any]]:
    """Export only scored runs, selecting public fields rather than whole configs."""
    local = _read_optional(manifest)
    assignment = local.get("assignment", {})
    outcomes = local.get("outcomes", {})
    result = local.get("result", {})
    rows = []
    for run in assignment.get("runs", []):
        side = run["side"]
        outcome = outcomes.get(side)
        if outcome is None:
            continue
        run_id = run["id"]
        if (
            not isinstance(run_id, str)
            or Path(run_id).name != run_id
            or run_id in {".", ".."}
            or "\\" in run_id
        ):
            raise ValueError("Invalid run ID in Excel source")
        root = manifest.parent / run_id
        score = outcome.get("partial_success")
        if (
            not isinstance(score, (int, float))
            or isinstance(score, bool)
            or not math.isfinite(score)
            or not 0 <= score <= 1
        ):
            raise ValueError("Saved partial success must be a finite score in [0, 1]")
        model = _read_optional(root / "policy.json").get("model", {})
        context = _read_optional(root / "recording-context.json")
        native = _read_optional(root / "inspect-result.json")
        stats = native.get("stats", {})
        dataset = root / "lerobot"
        replay = {
            title: dataset / f"videos/observation.images.{camera}/chunk-000/file-000.mp4"
            for title, camera in (
                ("Head video", "head_image"),
                ("Left video", "left_image"),
                ("Right video", "right_image"),
            )
        }
        mode = context.get("native_control_mode", "")
        rows.append(
            {
                "Router": local.get("excel_router", ""),
                "Evaluation ID": assignment["id"],
                "Run ID": run_id,
                "Track": assignment.get("track", ""),
                "Side": side,
                "Instruction": assignment.get("task", {}).get("instruction", ""),
                "Scene": assignment.get("scene", ""),
                "Evaluator": assignment.get("evaluator", ""),
                "Institution": assignment.get("institution", ""),
                "Model": model.get("name", ""),
                "Model URL": model.get("url", ""),
                "Revision": model.get("revision", ""),
                "Native control": "EEF / IK" if mode == "eef_abs_pose" else mode,
                "Success": outcome.get("success"),
                "Score (0-1)": score,
                "Score (%)": score * 100,
                "Preference": result.get("preference"),
                "Difficulty": assignment.get(
                    "difficulty", assignment.get("task", {}).get("difficulty")
                ),
                "Feedback": result.get("feedback", ""),
                "Test": assignment.get("test", False),
                "Result submitted": bool(local.get("submitted", False)),
                "Dataset uploaded": (
                    bool(local.get("datasets_uploaded", {}).get(run_id, False))
                    if assignment.get("recording", True)
                    else "disabled"
                ),
                "Steps": stats.get("total_steps"),
                "Duration (s)": stats.get("duration_s"),
                "Recorded at": stats.get("started_at", native.get("eval", {}).get("created", "")),
                "Inspect log": str(root / "inspect-result.json") if native else "",
                "Dataset folder": str(dataset) if dataset.is_dir() else "",
                **{title: str(path) if path.is_file() else "" for title, path in replay.items()},
            }
        )
    return rows


@contextmanager
def _lock(path: Path) -> Iterator[None]:
    # Spreadsheet editor locks are separate from our exporter lock. Never
    # replace an open workbook whose in-memory edits could overwrite our rows.
    if any(
        marker.exists()
        for marker in (
            path.parent / f".~lock.{path.name}#",
            path.parent / f"~${path.name}",
        )
    ):
        raise WorkbookBusy(f"Close {path.name} in Excel/LibreOffice before retrying")
    with path.with_suffix(path.suffix + ".colosseum.lock").open("a") as lock_file:
        try:
            import fcntl
        except ImportError:
            fcntl = None  # type: ignore[assignment]
        if fcntl is not None:
            try:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError as exc:
                raise WorkbookBusy("Another Colosseum Excel export is running") from exc
        try:
            yield
        finally:
            if fcntl is not None:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def _temp_path(parent: Path) -> Path:
    with tempfile.NamedTemporaryFile(
        prefix=".colosseum-excel-", suffix=".xlsx", dir=parent, delete=False
    ) as source:
        return Path(source.name)


def update_workbook(target: str | Path, manifests: list[Path]) -> int:
    """Atomically upsert by Router/evaluation/run ID; preserve all unrelated sheets."""
    path = validate_target(target)
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = [row for manifest in manifests for row in evaluation_rows(manifest)]
    if not rows:
        return 0
    path.parent.mkdir(parents=True, exist_ok=True)
    with _lock(path):
        existed = path.exists()
        book = load_workbook(path) if existed else Workbook()
        temp: Path | None = None
        try:
            new_sheet = SHEET not in book.sheetnames
            if new_sheet:
                sheet = book.create_sheet(SHEET)
                if not existed:
                    book.remove(book["Sheet"])
                sheet.append(HEADERS)
            else:
                sheet = book[SHEET]
                if tuple(c.value for c in sheet[1][: len(HEADERS)]) != HEADERS:
                    raise ValueError(f"Existing '{SHEET}' worksheet has incompatible headers")
            keys = {}
            for index in range(2, sheet.max_row + 1):
                key = tuple(sheet.cell(index, column).value for column in (1, 2, 3))
                if key in keys:
                    raise ValueError("Existing Colosseum worksheet has duplicate run rows")
                keys[key] = index
            for row in rows:
                key = tuple(row[header] for header in HEADERS[:3])
                index = keys.get(key, sheet.max_row + 1)
                keys[key] = index
                for column, header in enumerate(HEADERS, 1):
                    value = row.get(header)
                    if isinstance(value, (dict, list)):
                        value = json.dumps(value, ensure_ascii=False)
                    cell = sheet.cell(index, column, value)
                    # Assigning None through sheet.cell() does not clear a stale
                    # value; explicit assignment also keeps resume updates exact.
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = (
                            "s"  # Tasks/feedback are text, even if they start with '='.
                        )
                    cell.hyperlink = None
                    if header in _FILE_COLUMNS and value:
                        cell.hyperlink = os.path.relpath(value, path.parent)
                        cell.font = Font(color="2563EB", underline="single")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if header in {"Score (0-1)", "Score (%)", "Duration (s)"}:
                        cell.number_format = "0.00"
            for cell in sheet[1][: len(HEADERS)]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F2937")
            sheet.freeze_panes = "F2"
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
            for column, header in enumerate(HEADERS, 1):
                sheet.column_dimensions[get_column_letter(column)].width = (
                    45
                    if header in {"Instruction", "Feedback", *_FILE_COLUMNS}
                    else 30
                    if header in {"Scene", "Model URL", "Recorded at"}
                    else 20
                )
            if existed and new_sheet:
                backup = path.with_name(path.stem + ".before-colosseum.xlsx")
                if not backup.exists():
                    temp = _temp_path(path.parent)
                    shutil.copy2(path, temp)
                    os.replace(temp, backup)
                    temp = None
            temp = _temp_path(path.parent)
            book.save(temp)
            with temp.open("rb") as source:
                os.fsync(source.fileno())
            if existed:
                os.chmod(temp, path.stat().st_mode)
            os.replace(temp, path)
            temp = None
        finally:
            book.close()
            if temp is not None:
                temp.unlink(missing_ok=True)
    return len(rows)
